"""Rerunnable. Builds 07_revision3/EPCS_R3_MASTER_Results.xlsx plus a timestamped copy in
04_results/master_snapshots/. Run it after EVERY evaluation (it appends to evaluation_log.json).
READ-ONLY on run folders. CPU-light: no GPU, single thread. All commentary is generated from data."""
import os
os.environ["CUDA_VISIBLE_DEVICES"] = ""
os.environ["OMP_NUM_THREADS"] = "1"
os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"
import csv, json, shutil, math
from datetime import datetime
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

R = Path(__file__).resolve().parents[1]           # 07_revision3
RES = R / "04_results"
SNAP = RES / "master_snapshots"
LOG = SNAP / "evaluation_log.json"
OUT = R / "EPCS_R3_MASTER_Results.xlsx"
OLD_CSV = R.parent / "OLD_ROUNDS_1_2" / "EPCS_Master_Results_ForPaper.csv"
METRICS = ["accuracy", "precision", "recall", "specificity", "dice", "iou"]
SKIP_TOP = {"analysis", "splits", "medsam_r3", "sam_cache", "paper_tables", "logs", "master_snapshots"}
NOW = datetime.now()
TS = NOW.strftime("%Y-%m-%d %H:%M:%S")

try:
    from scipy import stats as _st
    def tcrit(df): return float(_st.t.ppf(0.975, df))
except Exception:
    def tcrit(df): return 1.96


# ---------------------------------------------------------------- per-run table
def load_runs():
    rows = []
    for f in sorted(RES.rglob("summary.json")):
        rel = f.relative_to(RES).parts
        if any(s.startswith("_") for s in rel) or rel[0] in SKIP_TOP:
            continue
        try:
            j = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            continue
        if "test_mean" not in j or "model" not in j:
            continue
        r = dict(model=j["model"], label=j.get("model_label"), group=j.get("group"), dataset=j.get("dataset"),
                 fold=j.get("fold"), seed=j.get("seed"), best_epoch=j.get("best_epoch"))
        for m in METRICS:
            r["val_" + m] = j.get("validation_mean", {}).get(m)
        for m in METRICS:
            r["test_" + m] = j["test_mean"].get(m)
        r.update(params_total=j.get("parameters_total"), params_trainable=j.get("parameters_trainable"),
                 wall_seconds=j.get("wall_seconds"), n_test_images=j.get("n_test"))
        rows.append(r)
    return pd.DataFrame(rows)


def agg(df):
    out = []
    for (ds, md), g in df.groupby(["dataset", "model"]):
        r = dict(dataset=ds, model=md, label=g["label"].iloc[0], group=g["group"].iloc[0], n_runs=len(g),
                 n_seeds=g["seed"].nunique(), n_folds=g["fold"].nunique(),
                 params_total=g["params_total"].iloc[0], params_trainable=g["params_trainable"].iloc[0],
                 best_epoch_mean=g["best_epoch"].mean())
        n = len(g)
        for m in METRICS:
            v = g["test_" + m].astype(float)
            mean = v.mean(); sd = v.std(ddof=1) if n > 1 else np.nan
            h = tcrit(n - 1) * sd / math.sqrt(n) if n > 1 else np.nan
            r.update({f"test_{m}_mean": mean, f"test_{m}_sd": sd, f"test_{m}_ci_low": mean - h, f"test_{m}_ci_high": mean + h})
        for fo, gg in g.groupby("fold"):
            r[f"fold{fo}_test_dice_mean"] = gg["test_dice"].mean()
        out.append(r)
    a = pd.DataFrame(out)
    a["rank_by_test_dice"] = a.groupby("dataset")["test_dice_mean"].rank(ascending=False, method="min").astype(int)
    return a.sort_values(["dataset", "rank_by_test_dice"]).reset_index(drop=True)


def best_values(df):
    out = []
    for (ds, md), g in df.groupby(["dataset", "model"]):
        for m in METRICS:
            c = "test_" + m
            i = g[c].astype(float).idxmax()
            out.append(dict(dataset=ds, model=md, metric=m, best_single_run_value=g.loc[i, c], best_fold=g.loc[i, "fold"],
                            best_seed=g.loc[i, "seed"], best_epoch=g.loc[i, "best_epoch"], median=g[c].median(),
                            mean=g[c].mean(), worst=g[c].min(), n_runs=len(g)))
    return pd.DataFrame(out)


# ---------------------------------------------------------------- previous rounds
def read_old():
    """Parse the old master csv: Run A table, Run B table, MedSAM zero-shot table (all validation-set numbers)."""
    A, B, M = [], [], []
    if not OLD_CSV.exists():
        return A, B, M
    rows = list(csv.reader(open(OLD_CSV, encoding="utf-8-sig")))
    mode = None
    for r in rows:
        if not r or not any(r):
            mode = None if mode and len(r) == 0 else mode
            continue
        if r[0] == "Dataset" and len(r) > 2:
            mode = "A" if r[1] == "Model" else "B" if r[1] == "Seeds" else "M" if r[1] == "Accuracy" else None
            hdr = r
            continue
        if mode and r[0] in ("CryoNuSeg", "NuInsSeg"):
            d = dict(zip(hdr, r))
            if mode == "A": A.append(d)
            elif mode == "B": B.append(d)
            elif mode == "M": M.append(d)
    return A, B, M


MAP_A = {"U-Net": "unet", "U-Net++": "unetpp", "U-Net3+": "unet3plus", "eU-Net3+ (no SAM)": "eunet3plus",
         "SAM-guided baseline": "sam_unet", "EPCS-Net (proposed)": "epcs_net"}
COLS_OLD = {"accuracy": "Accuracy", "precision": "Precision", "recall": "Recall", "specificity": "Specificity",
            "dice": "Dice", "iou": "IoU"}


def comment(prev, new, n, newprot):
    if new is None or (isinstance(new, float) and np.isnan(new)):
        return "No new test-split result yet for this model and dataset; nothing to compare."
    d = new - prev
    if abs(d) < 0.005:
        s = f"Essentially unchanged (difference {d:+.3f})."
    elif d < 0:
        s = (f"Lower by {abs(d):.3f}, expected because previous scores were measured on the checkpoint-selection "
             f"(validation) set, which favours the selected checkpoint, whereas new scores are on the untouched test split.")
    else:
        s = (f"Higher by {d:.3f}; the two numbers use different splits and protocols, so this is a descriptive "
             f"difference only.")
    return s + f" New value is a mean over {n} run(s); no significance is claimed."


def vs_previous(avg, medsam_means):
    A, B, M = read_old()
    rows = []
    def newv(ds, md, m):
        s = avg[(avg.dataset == ds) & (avg.model == md)]
        if s.empty: return None, 0
        return float(s.iloc[0][f"test_{m}_mean"]), int(s.iloc[0]["n_runs"])
    for a in A:
        ds = a["Dataset"].lower(); md = MAP_A.get(a["Model"])
        for m in METRICS:
            try: p = float(a[COLS_OLD[m]])
            except Exception: continue
            nv, n = newv(ds, md, m)
            rows.append(dict(source="Run A", dataset=ds, model=md, metric=m, previous_value=p,
                             previous_protocol="validation-selected, single run (Run A, 1 seed)",
                             new_value=nv, new_protocol=f"test split, mean over {n} runs" if n else "not yet run",
                             delta=None if nv is None else nv - p, comment=comment(p, nv, n, "")))
    for b in B:
        ds = b["Dataset"].lower()
        for m in METRICS:
            try: p = float(b[COLS_OLD[m] + (" (mean)" if m in ("dice", "iou") else "")])
            except Exception: continue
            nv, n = newv(ds, "epcs_net", m)
            rows.append(dict(source="Run B", dataset=ds, model="epcs_net", metric=m, previous_value=p,
                             previous_protocol="validation-selected, 5-seed mean (Run B, same split as Run A)",
                             new_value=nv, new_protocol=f"test split, mean over {n} runs" if n else "not yet run",
                             delta=None if nv is None else nv - p, comment=comment(p, nv, n, "")))
    for mm in M:
        ds = mm["Dataset"].lower()
        for m in METRICS:
            try: p = float(mm[COLS_OLD[m]])
            except Exception: continue
            nv = medsam_means.get((ds, "zero_shot_wholebox", m))
            rows.append(dict(source="MedSAM old", dataset=ds, model="medsam_zero_shot", metric=m, previous_value=p,
                             previous_protocol="validation set, zero-shot whole-image box, single run",
                             new_value=nv, new_protocol="test split, zero-shot whole-image box, mean over folds" if nv is not None else "not yet run",
                             delta=None if nv is None else nv - p, comment=comment(p, nv, 3 if ds == "cryonuseg" else 1, "")))
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- ablation / medsam / significance
ABL = ["abl0_nosam_epcsbackbone", "abl1_nosam", "abl2_bottleneck", "abl3_pspi", "abl4_bdsg", "abl5_dksr", "abl6_ftas",
       "ablA_acfb_nopspi", "ablB_random_embed", "ablC_pspi_concat", "ablD_gn_encoder"]


def ablation(avg):
    pt = RES / "analysis" / "paired_significance_tests.csv"
    ptd = pd.read_csv(pt) if pt.exists() else pd.DataFrame()
    if not ptd.empty:
        ptd = ptd[(ptd.family == "ablation") & (ptd.metric == "dice")]
    rows = []
    for ds in sorted(avg.dataset.unique()):
        full = avg[(avg.dataset == ds) & (avg.model == "epcs_net")]
        fm = float(full.iloc[0]["test_dice_mean"]) if not full.empty else np.nan
        for md in ["epcs_net"] + ABL:
            s = avg[(avg.dataset == ds) & (avg.model == md)]
            if s.empty:
                rows.append(dict(dataset=ds, model=md, n_runs=0, comment="No completed runs yet.")); continue
            s = s.iloc[0]; mean, sd = s["test_dice_mean"], s["test_dice_sd"]
            r = dict(dataset=ds, model=md, label=s["label"], n_runs=s["n_runs"], test_dice_mean=mean, test_dice_sd=sd,
                     mean_pm_sd="%.4f ± %s" % (mean, "n/a" if np.isnan(sd) else "%.4f" % sd),
                     test_iou_mean=s["test_iou_mean"], delta_dice_vs_full=mean - fm if md != "epcs_net" else 0.0)
            if md == "epcs_net":
                r["comment"] = "Full model (reference row)."
            else:
                p = ptd[(ptd.dataset == ds) & (ptd.baseline == md)] if not ptd.empty else pd.DataFrame()
                c = f"Dice {r['delta_dice_vs_full']:+.4f} vs the full model (mean over {int(s['n_runs'])} run(s))."
                if not p.empty:
                    p = p.iloc[0]
                    r.update(paired_n=p.n_pairs, paired_mean_diff_full_minus_variant=p.mean_difference,
                             ci95_low=p.ci95_low, ci95_high=p.ci95_high, t_p=p.t_p_value, t_p_holm=p.t_p_holm,
                             wilcoxon_p_holm=p.wilcoxon_p_holm)
                    ex = "excludes" if (p.ci95_low > 0 or p.ci95_high < 0) else "includes"
                    c += (f" Paired test on {int(p.n_pairs)} matched pairs: full minus variant {p.mean_difference:+.4f}, 95% CI "
                          f"[{p.ci95_low:+.4f}, {p.ci95_high:+.4f}] {ex} zero; Holm-adjusted t p={p.t_p_holm:.3f}. "
                          "Seeds within a fold share test images, so p-values are approximate; read as descriptive.")
                else:
                    c += " No paired test row available."
                r["comment"] = c
            rows.append(r)
    d = pd.DataFrame(rows)
    return d[[c for c in d.columns if c != 'comment'] + ['comment']]


def medsam():
    f = RES / "medsam_r3" / "summary_all.json"
    rows, means = [], {}
    if not f.exists(): return pd.DataFrame(), means
    for j in json.loads(f.read_text(encoding="utf-8")):
        if j.get("dry_run"): continue
        r = dict(mode=j["mode"], dataset_split=j["dataset"], n_test_images=j.get("n_test_images"), protocol=j.get("protocol"))
        for m in METRICS:
            r[m] = j["metric_mean"].get(m); r[m + "_sd_images"] = j.get("metric_sample_sd", {}).get(m)
        rows.append(r)
    df = pd.DataFrame(rows)
    if df.empty: return df, means
    df["dataset"] = df.dataset_split.str.replace(r"_fold\d+$", "", regex=True)
    for (ds, mo), g in df.groupby(["dataset", "mode"]):
        for m in METRICS: means[(ds, mo, m)] = float(g[m].mean())
    role = {"zero_shot_wholebox": "sanity check (not a competing method)", "oracle_boxes": "upper bound (ground-truth boxes, not deployable)"}
    df["role"] = df["mode"].map(role)
    mean_rows = []
    for (ds, mo), g in df.groupby(["dataset", "mode"]):
        r = dict(mode=mo, dataset_split=ds + " (mean over splits)", n_test_images=g.n_test_images.sum(), protocol=g.protocol.iloc[0], dataset=ds, role=role.get(mo))
        for m in METRICS: r[m] = g[m].mean()
        mean_rows.append(r)
    df = pd.concat([df, pd.DataFrame(mean_rows)], ignore_index=True)
    cols = ["dataset", "mode", "role", "dataset_split", "n_test_images"] + METRICS + [m + "_sd_images" for m in METRICS] + ["protocol"]
    return df[cols].sort_values(["dataset", "mode", "dataset_split"]).reset_index(drop=True), means


# ---------------------------------------------------------------- evaluation log
def update_log(runs, avg):
    SNAP.mkdir(parents=True, exist_ok=True)
    log = json.loads(LOG.read_text(encoding="utf-8")) if LOG.exists() else []
    per_ds = {k: int(v) for k, v in runs.groupby("dataset").size().items()}
    head = {}
    for ds, g in avg.groupby("dataset"):
        e = g[g.model == "epcs_net"]
        b = g[(g.group == "baseline")].sort_values("test_dice_mean", ascending=False)
        head[ds] = dict(epcs_net_test_dice=None if e.empty else float(e.iloc[0].test_dice_mean),
                        epcs_net_n_runs=None if e.empty else int(e.iloc[0].n_runs),
                        best_baseline=None if b.empty else b.iloc[0].model,
                        best_baseline_test_dice=None if b.empty else float(b.iloc[0].test_dice_mean))
    entry = dict(timestamp=TS, runs_total=int(len(runs)), runs_per_dataset=per_ds, headline=head)
    if log:
        p = log[-1]; parts = []
        dr = entry["runs_total"] - p.get("runs_total", 0)
        parts.append(f"{dr:+d} completed runs since last evaluation ({p.get('runs_total', 0)} -> {entry['runs_total']}).")
        for ds, h in head.items():
            ph = p.get("headline", {}).get(ds)
            if ph is None:
                parts.append(f"{ds}: first evaluation of this dataset; EPCS-Net test Dice {h['epcs_net_test_dice']}.") if h["epcs_net_test_dice"] is not None else None
            elif h["epcs_net_test_dice"] is not None and ph.get("epcs_net_test_dice") is not None:
                d = h["epcs_net_test_dice"] - ph["epcs_net_test_dice"]
                bd = (h["best_baseline_test_dice"] - ph["best_baseline_test_dice"]) if h["best_baseline_test_dice"] is not None and ph.get("best_baseline_test_dice") is not None else None
                s = f"{ds}: EPCS-Net test Dice {h['epcs_net_test_dice']:.4f} ({d:+.4f}, runs {ph.get('epcs_net_n_runs')} -> {h['epcs_net_n_runs']})"
                if bd is not None: s += f"; best baseline {h['best_baseline']} {h['best_baseline_test_dice']:.4f} ({bd:+.4f})"
                parts.append(s + ".")
        for ds in per_ds:
            if ds not in p.get("runs_per_dataset", {}):
                parts.append(f"{ds}: new dataset appeared in the evaluation ({per_ds[ds]} runs).")
        if dr == 0: parts.append("No new runs; numbers unchanged.")
        entry["comment"] = " ".join(x for x in parts if x)
    else:
        entry["comment"] = "First logged evaluation; no previous entry to compare with."
    log.append(entry)
    LOG.write_text(json.dumps(log, indent=2), encoding="utf-8")
    rows = []
    for e in log:
        r = dict(timestamp=e["timestamp"], runs_total=e["runs_total"], runs_per_dataset="; ".join(f"{k}: {v}" for k, v in e["runs_per_dataset"].items()))
        for ds, h in e["headline"].items():
            r[f"{ds}_epcs_net_test_dice"] = h["epcs_net_test_dice"]
            r[f"{ds}_best_baseline"] = h["best_baseline"]
            r[f"{ds}_best_baseline_test_dice"] = h["best_baseline_test_dice"]
        r["comment"] = e["comment"]
        rows.append(r)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- main
def style(ws, header_row=1):
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for i, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) for c in list(col)[:200] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 70)
    for row in ws.iter_rows(min_row=header_row + 1):
        for c in row:
            if isinstance(c.value, float): c.number_format = "0.0000"


def main():
    runs = load_runs()
    if runs.empty:
        raise SystemExit("no completed runs found")
    runs = runs.sort_values(["dataset", "group", "model", "fold", "seed"]).reset_index(drop=True)
    per_run = runs[["model", "label", "group", "dataset", "fold", "seed", "best_epoch"] + ["val_" + m for m in METRICS]
                   + ["test_" + m for m in METRICS] + ["params_total", "params_trainable", "wall_seconds"]]
    avg = agg(runs)
    best = best_values(runs)
    ms, ms_means = medsam()
    vp = vs_previous(avg, ms_means)
    abl = ablation(avg)
    logdf = update_log(runs, avg)

    sigdir = RES / "analysis"
    sig_files = [("paired_significance_tests.csv", "Paired tests: proposed vs baseline / ablation (one pair per matched fold x seed run)"),
                 ("fold_averaged_tests.csv", "Fold-averaged tests (seeds averaged within fold, n = folds)"),
                 ("bootstrap_per_image.csv", "Bootstrap over test images")]
    readme = [
        ["EPCS-Net Round 3 master results workbook"], [f"Built {TS} by 06_scripts/build_master_workbook.py (rerun after every evaluation)."],
        [f"Completed runs read: {len(runs)} (summary.json files, test split; paths with a leading underscore and non-run folders skipped)."],
        [""], ["Sheets"],
        ["Master_PerRun: every completed run with validation and test values of six metrics, parameters, wall seconds."],
        ["Master_Averages: per dataset x model, mean, sample SD, 95% t-CI (df = n-1) over runs; per-fold mean Dice; rank by mean test Dice."],
        ["Best_Values: best single run per metric (with fold, seed, epoch) plus median/mean/worst."],
        ["Vs_Previous: old Run A / Run B / MedSAM numbers (validation-selected) against new test-split means. Different protocols; descriptive only."],
        ["Ablation: 11 variants + full model with delta vs full and commentary from paired tests."],
        ["MedSAM: zero-shot whole-box as sanity check, oracle GT boxes as upper bound (not competing methods)."],
        ["Significance: copies of the stats csvs."],
        ["Evaluation_Log: append-only history (04_results/master_snapshots/evaluation_log.json)."],
        [""], ["Caveats"],
        ["Test metrics are the mean over test images per run. Seeds within a fold share the same test images, so CIs and p-values treat runs as independent only approximately."],
        ["Previous numbers were measured on the validation set that also chose the checkpoint; the new numbers are on the untouched test split. Differences reflect a protocol change as well as any model effect."],
        ["All COMMENT strings are generated by code from the numbers shown; no significance is claimed anywhere."],
    ]

    sheets = [("Master_PerRun", per_run), ("Master_Averages", avg), ("Best_Values", best), ("Vs_Previous", vp),
              ("Ablation", abl), ("MedSAM", ms)]
    snap_name = SNAP / f"EPCS_R3_MASTER_{NOW.strftime('%Y%m%d_%H%M')}.xlsx"
    tmp = SNAP / "_tmp_master.xlsx"
    counts = {}
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        for n, d in sheets:
            d.to_excel(xw, sheet_name=n, index=False); counts[n] = len(d); style(xw.sheets[n])
        # Significance
        ws = xw.book.create_sheet("Significance"); r = 1
        ws.cell(r, 1, "WARNING: seeds within a fold share the same test images, so runs are not independent. P-values and CIs here are approximate, "
                      "small-n descriptive aids and do not establish significance. Fold-averaged tests have very low power (n = folds).").font = Font(bold=True, color="C00000")
        r += 2; nrows = 0
        for fn, title in sig_files:
            p = sigdir / fn
            ws.cell(r, 1, title + f"  [{fn}]").font = Font(bold=True); r += 1
            if not p.exists(): ws.cell(r, 1, "file not found"); r += 2; continue
            d = pd.read_csv(p)
            for j, c in enumerate(d.columns, 1):
                cell = ws.cell(r, j, c); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="305496")
            r += 1
            for row in d.itertuples(index=False):
                for j, v in enumerate(row, 1):
                    if isinstance(v, (np.generic,)): v = v.item()
                    if isinstance(v, float) and np.isnan(v): v = None
                    ws.cell(r, j, v)
                r += 1; nrows += 1
            r += 1
        counts["Significance"] = nrows
        for i in range(1, 30): ws.column_dimensions[get_column_letter(i)].width = 16
        logdf.to_excel(xw, sheet_name="Evaluation_Log", index=False); counts["Evaluation_Log"] = len(logdf); style(xw.sheets["Evaluation_Log"])
        ws = xw.book.create_sheet("Readme")
        for i, row in enumerate(readme, 1): ws.cell(i, 1, row[0])
        ws["A1"].font = Font(bold=True, size=14); ws.column_dimensions["A"].width = 150; counts["Readme"] = len(readme)
    shutil.copy(tmp, snap_name)
    try:
        shutil.copy(tmp, OUT)
    except PermissionError:
        print("WARNING: could not overwrite", OUT, "(open in Excel?); snapshot written only")
    tmp.unlink()
    print("runs:", len(runs), "snapshot:", snap_name)
    for k, v in counts.items(): print(f"  {k}: {v} rows")


if __name__ == "__main__":
    main()
