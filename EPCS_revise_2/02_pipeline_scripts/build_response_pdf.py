"""Build the point-by-point response PDF (and .docx if python-docx exists)
from 05_manuscript/RESPONSE_DRAFT_R3.md.  Output name is ..._DRAFT.pdf while any
[[RESULT: ...]] placeholder remains; ..._FINAL.pdf only at zero placeholders.
Run:  python build_response_pdf.py   (CPU only, a few seconds)"""
import os, re
os.environ.setdefault("CUDA_VISIBLE_DEVICES", "")
os.environ.setdefault("OMP_NUM_THREADS", "1")
from xml.sax.saxutils import escape
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                Table, TableStyle, HRFlowable, PageBreak)

HERE = os.path.dirname(os.path.abspath(__file__))
MS = os.path.normpath(os.path.join(HERE, "..", "05_manuscript"))
SRC = os.path.join(MS, "RESPONSE_DRAFT_R3.md")
OUT = os.path.join(MS, "response")
TITLE = "EPCS-Net: SAM-Guided Multi-Scale Nuclei Segmentation"
SUBID = "75a92866-8513-48b6-9941-b16b3a741de2"
JOURNAL = "Scientific Reports"
RES = re.compile(r"\[\[RESULT:.*?\]\]", re.S)

text = open(SRC, encoding="utf-8").read()
nph = len(RES.findall(text))
suffix = "DRAFT" if nph else "FINAL"
if nph:
    print("REFUSING FINAL: %d [[RESULT]] placeholder(s) remain; writing DRAFT only." % nph)
os.makedirs(OUT, exist_ok=True)
pdf_path = os.path.join(OUT, "EPCS_R3_Response_to_Reviewers_%s.pdf" % suffix)
docx_path = pdf_path[:-4] + ".docx"
for s in ("DRAFT", "FINAL"):  # remove stale file of the other kind
    if s != suffix:
        for ext in (".pdf", ".docx"):
            p = os.path.join(OUT, "EPCS_R3_Response_to_Reviewers_%s%s" % (s, ext))
            if os.path.exists(p):
                os.remove(p)


# ---- LaTeX label -> number resolution from the manuscript tex -------------------------
TEX = os.path.join(MS, "EPCS_revise3.tex")
LABELS = {}


def _build_labels():
    if not os.path.exists(TEX):
        return
    c = [0, 0, 0]
    lvl_now = 0
    cnt = {"tab": 0, "fig": 0, "eq": 0}
    for ln in open(TEX, encoding="utf-8", errors="replace"):
        m = re.match(r"\s*\\((?:sub)*)section\{", ln)
        if m:
            lvl_now = len(m.group(1)) // 3
            c[lvl_now] += 1
            for j in range(lvl_now + 1, 3):
                c[j] = 0
        if re.match(r"\s*\\begin\{table\*?\}", ln):
            cnt["tab"] += 1
        if re.match(r"\s*\\begin\{figure\*?\}", ln):
            cnt["fig"] += 1
        for lab in re.findall(r"\\label\{([^}]+)\}", ln):
            if lab.startswith("sec:"):
                LABELS[lab] = ".".join(str(x) for x in c[:lvl_now + 1])
            elif lab.startswith("tab:"):
                LABELS[lab] = str(cnt["tab"])
            elif lab.startswith("fig:"):
                LABELS[lab] = str(cnt["fig"])
            elif lab.startswith("eq:"):
                cnt["eq"] += 1
                LABELS[lab] = str(cnt["eq"])


_build_labels()
UNRESOLVED = set()


def resolve_refs(s):
    def one(m):
        lab = m.group(1)
        if lab in LABELS:
            return LABELS[lab]
        UNRESOLVED.add(lab)
        return lab.split(":", 1)[-1].replace("_", " ")
    s = re.sub(r"~?\\(?:eq)?ref\{([^}]+)\}", lambda m: (" " if m.group(0).startswith("~") else "") + one(m), s)
    s = re.sub(r"\\cite\{[^}]*\}", "", s)
    return s


# ---- verbatim reviewer comments ---------------------------------------------------------
VERB = os.path.join(OUT, "REVIEWER_COMMENTS_VERBATIM.md")
VCOM = {}      # parent id -> verbatim text (paragraphs joined by newline)
VUSED = set()


def _load_verbatim():
    cur, sect = None, None
    for raw in open(VERB, encoding="utf-8").read().split("\n"):
        ln = raw.rstrip()
        if ln.startswith("# "):
            continue
        if ln.startswith("## "):
            sect = ln[3:]
            cur = None
            continue
        if not ln.strip():
            continue
        if sect and sect.startswith("Editor"):
            if ln.startswith("Code policy:"):
                cur = "Ed-5"; VCOM[cur] = ln[len("Code policy:"):].strip()
            elif ln.startswith("Submission requirements:"):
                cur = "Ed-6"; VCOM[cur] = ln[len("Submission requirements:"):].strip()
            elif re.match(r"\d\. ", ln):
                cur = "Ed-" + ln[0]; VCOM[cur] = ln[3:].strip()
            elif ln.startswith('"According'):
                cur = "Ed-0b"; VCOM[cur] = ln
            elif cur is None or cur == "Ed-0":
                cur = "Ed-0"; VCOM[cur] = (VCOM.get(cur, "") + " " + ln).strip()
            else:
                VCOM[cur] += "\n" + ln
            continue
        m = re.match(r"^(R\d-[Mm]?\d+)\.\s*(.*)", ln)
        m2 = re.match(r"^(\d)\. (.*)", ln) if sect == "Reviewer 2" else None
        if m:
            cur = m.group(1); VCOM[cur] = m.group(2)
        elif m2:
            cur = "R2-" + m2.group(1); VCOM[cur] = m2.group(2)
        elif cur and not ln.startswith(("Major comments", "Minor comments")):
            VCOM[cur] += "\n" + ln


_load_verbatim()


def parent_id(cid):
    cid = cid.strip().rstrip(".")
    m = re.match(r"^(Ed-\d+|R\d-[Mm]?\d+)([a-z]?)$", cid)
    return (m.group(1), bool(m.group(2))) if m else (cid, False)


def norm(s):
    s = s.replace("+/-", "\u00b1")
    s = re.sub(r"(?<=\d) x (?=\d)", " \u00d7 ", s)
    s = s.replace(">=", "\u2265").replace("<=", "\u2264")
    return s


FD = "C:/Windows/Fonts/"
pdfmetrics.registerFont(TTFont("Body", FD + "arial.ttf"))
pdfmetrics.registerFont(TTFont("Body-B", FD + "arialbd.ttf"))
pdfmetrics.registerFont(TTFont("Body-I", FD + "ariali.ttf"))
pdfmetrics.registerFont(TTFont("Body-BI", FD + "arialbi.ttf"))
pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-B", italic="Body-I", boldItalic="Body-BI")
pdfmetrics.registerFont(TTFont("Mono", FD + "consola.ttf"))

NAVY = colors.HexColor("#1f3a5f")
base = ParagraphStyle("b", fontName="Body", fontSize=9.5, leading=13.2, spaceAfter=5)
st = dict(
    base=base,
    h2=ParagraphStyle("h2", parent=base, fontName="Body-B", fontSize=13, leading=16, textColor=NAVY, spaceBefore=10, spaceAfter=6),
    cbox=ParagraphStyle("c", parent=base, fontSize=9.3, leading=12.5, spaceAfter=0),
    bullet=ParagraphStyle("bl", parent=base, leftIndent=14, bulletIndent=3),
    cell=ParagraphStyle("cell", parent=base, fontSize=8, leading=10, spaceAfter=0),
    cellh=ParagraphStyle("cellh", parent=base, fontName="Body-B", fontSize=8, leading=10, spaceAfter=0, textColor=colors.white),
    title=ParagraphStyle("t", parent=base, fontName="Body-B", fontSize=20, leading=25, textColor=NAVY, spaceAfter=4),
    meta=ParagraphStyle("m", parent=base, fontSize=10.5, leading=15),
)


def inline(s):
    s = resolve_refs(s)
    out, pos = [], 0

    def fmt(seg):
        seg = escape(norm(seg))
        seg = re.sub(r"`([^`]+)`", lambda m: '<font name="Mono" size="8.3">%s</font>' % m.group(1), seg)
        seg = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", seg)
        return seg
    for m in RES.finditer(s):
        out.append(fmt(s[pos:m.start()]))
        out.append('<font backColor="#FFEB3B" color="#7a0000"><b>%s</b></font>'
                   % escape(norm(" ".join(m.group(0).split()))))
        pos = m.end()
    out.append(fmt(s[pos:]))
    return "".join(out)


lines = text.split("\n")
blocks = []
i = 0
while i < len(lines):
    ln = lines[i]
    if not ln.strip():
        i += 1
        continue
    if ln.startswith("# "):
        blocks.append(("h1", ln[2:].strip()))
    elif ln.startswith("## "):
        blocks.append(("h2", ln[3:].strip()))
    elif ln.strip() == "---":
        blocks.append(("hr", None))
    elif ln.lstrip().startswith("|"):
        rows = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        blocks.append(("table", rows))
        continue
    elif re.match(r"^\*\*(R\d|Ed)-.*\*\*\s*$", ln):
        blocks.append(("comment", ln.strip()[2:-2]))
    elif re.match(r"^\s*[-*] ", ln):
        blocks.append(("bullet", re.sub(r"^\s*[-*] ", "", ln)))
    elif re.match(r"^\d+\. ", ln):
        n, t = ln.split(". ", 1)
        blocks.append(("num", (n, t)))
    else:
        blocks.append(("p", ln.strip()))
    i += 1

ADD = [
    "<b>Held-out test splits.</b> Train/validation/test replaces the old validation-as-reporting protocol; checkpoints are chosen on validation only; CryoNuSeg uses a tissue-stratified three-fold scheme, NuInsSeg an organ-stratified split (R3-M1, R4-1, R2-2).",
    "<b>Matched five-seed baselines.</b> All eight models, baselines included, are trained with the same five seeds on identical splits (15 matched CryoNuSeg runs per model), with seed-level and conservative fold-level paired tests (R2-1, R3-M3, R1-1).",
    "<b>Five-seed ablation with controls.</b> The ablation is rerun with five seeds (CryoNuSeg fold 1), including ACFB-without-PSPI, random-embedding and GroupNorm-encoder controls; the earlier claim that every component earns its place is withdrawn (R1-2, R3-M4).",
    "<b>Reference-21 benchmark.</b> A verified re-implementation of the fusion block of reference 21 is trained under the common protocol at the bottleneck and at four scales (R4-2).",
    "<b>MedSAM sanity and oracle checks.</b> Zero-shot and ground-truth-box (oracle) MedSAM runs, reported as a sanity check and an upper bound, not as a fair competitor (R1-4).",
    "<b>OSPT sampling analysis.</b> Table quantifying the sampling behaviour behind the &quot;balanced within every epoch&quot; statement (R1-3).",
    "<b>Latency re-measurement.</b> The non-comparable published latencies are explained; a controlled re-measurement (GPU model, warm-up, iterations) is specified (R3-m2, R4-4c).",
    "<b>Grad-CAM failure panel.</b> Lowest-Dice test cases chosen by metric and shown beside best cases (R2-3).",
]


def footer(c, d):
    c.saveState()
    c.setFont("Body", 8)
    c.setFillColor(colors.HexColor("#555555"))
    c.drawString(18 * mm, 10 * mm, "Response to Reviewers, Round 3 | Submission %s" % SUBID)
    c.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % d.page)
    if nph:
        c.setFillColor(colors.HexColor("#b00020"))
        c.setFont("Body-B", 8)
        c.drawCentredString(A4[0] / 2, A4[1] - 9 * mm, "DRAFT - NOT FOR SUBMISSION - %d unresolved placeholder(s)" % nph)
    c.restoreState()


doc = BaseDocTemplate(pdf_path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=18 * mm, bottomMargin=17 * mm,
                      title="Response to Reviewers - " + TITLE, author="Authors")
W = A4[0] - 36 * mm
doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(18 * mm, 17 * mm, W, A4[1] - 35 * mm, id="f",
                      leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)], onPage=footer)])

meta = [["Manuscript", TITLE], ["Submission ID", SUBID], ["Journal", JOURNAL]]
S = [Paragraph("Point-by-Point Response to the Editor and Reviewers", st["title"]),
     Paragraph("Revision round 3" + (" (DRAFT)" if nph else ""), st["meta"]), Spacer(1, 4)]
mt = Table([[Paragraph("<b>%s</b>" % a, st["cell"]), Paragraph(escape(b), st["cell"])] for a, b in meta],
           colWidths=[32 * mm, W - 32 * mm])
mt.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, NAVY), ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf0f8")),
                        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
S += [mt, Spacer(1, 8), Paragraph("Summary of additional experiments", st["h2"])]
for a in ADD:
    S.append(Paragraph(a, st["bullet"], bulletText="\u2022"))
S.append(Paragraph("Each shaded box shows our short heading in bold followed by the comment verbatim; our responses follow each box. "
                   "Text on a yellow background marks a result that is still pending.", st["base"]))
if nph:
    S.append(Paragraph('<font backColor="#FFEB3B" color="#7a0000"><b>This is a DRAFT: %d placeholder(s) remain.</b></font>' % nph, st["base"]))
S.append(PageBreak())


def table_flow(rows):
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    w = [min(max(max(len(r[j]) for r in rows), 6), 60) for j in range(n)]
    tot = sum(w)
    cw = [W * x / tot for x in w]
    data = [[Paragraph(inline(c), st["cellh"] if ri == 0 else st["cell"]) for c in r] for ri, r in enumerate(rows)]
    t = Table(data, colWidths=cw, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6fa")]),
                           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    return t


SEEN, UNMATCHED = set(), []
for k, p in blocks:
    if k == "h1":
        continue
    elif k == "h2":
        S.append(Paragraph(inline(p), st["h2"]))
    elif k == "hr":
        S.append(HRFlowable(width="100%", thickness=0.5, color=colors.lightgrey, spaceBefore=2, spaceAfter=4))
    elif k == "comment":
        m = re.match(r"(\S+?)\.?\s+(.*)", p, re.S)
        cid, body = (m.group(1), m.group(2)) if m else ("", p)
        cid = cid.rstrip(".")
        par, is_sub = parent_id(cid)
        verb = None if par in SEEN else VCOM.get(par)
        SEEN.add(par)
        if verb:
            VUSED.add(par)
        elif par not in VCOM:
            UNMATCHED.append(cid)
        rows = [[Paragraph("<b>%s. %s</b>" % (escape(cid), inline(body)), st["cbox"])]]
        pre = []
        if par == "Ed-1":
            for g in ("Ed-0", "Ed-0b"):
                if g in VCOM:
                    VUSED.add(g)
                    pre.append(Paragraph("<font size=8 color='#555555'>Editor, %s (verbatim)</font><br/>%s" % (
                        "general comment" if g == "Ed-0" else "decision note", inline(VCOM[g])), st["cbox"]))
        if verb:
            pre.append(Paragraph("<font size=8 color='#555555'>Reviewer comment, verbatim</font>", st["cbox"]))
            for part in verb.split("\n"):
                pre.append(Paragraph(inline(part), st["cbox"]))
        rows += [[x] for x in pre]
        box = Table(rows, colWidths=[W], splitByRow=1)
        box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#e9edf3")),
                                 ("LINEBEFORE", (0, 0), (0, -1), 3, NAVY),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                                 ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        S += [Spacer(1, 4), box, Spacer(1, 4)]
    elif k == "table":
        S += [table_flow(p), Spacer(1, 6)]
    elif k == "bullet":
        S.append(Paragraph(inline(p), st["bullet"], bulletText="\u2022"))
    elif k == "num":
        S.append(Paragraph(inline(p[1]), st["bullet"], bulletText=p[0] + "."))
    else:
        s = inline(p)
        s = re.sub(r"^((?:Response|Location)[^:<]*:)", r"<b>\1</b>", s, count=1)
        S.append(Paragraph(s, st["base"]))
doc.build(S)
print("PDF:", pdf_path)

try:
    import docx
    from docx.shared import Pt
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    D = docx.Document()
    D.styles["Normal"].font.name = "Arial"
    D.styles["Normal"].font.size = Pt(10)

    def runs(par, s):
        s = resolve_refs(s)
        def seg(t, hl=False):
            for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", norm(t)):
                if not part:
                    continue
                b = part.startswith("**") and part.endswith("**") and len(part) > 4
                c = part.startswith("`") and part.endswith("`")
                r = par.add_run(part[2:-2] if b else part[1:-1] if c else part)
                r.bold = b or hl
                if c:
                    r.font.name = "Consolas"
                if hl:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = 0
        for m in RES.finditer(s):
            seg(s[pos:m.start()])
            seg(" ".join(m.group(0).split()), True)
            pos = m.end()
        seg(s[pos:])

    def shade(cell, hexc):
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), hexc)
        cell._tc.get_or_add_tcPr().append(sh)
    D.add_heading("Point-by-Point Response to the Editor and Reviewers", 0)
    for a, b in meta:
        par = D.add_paragraph()
        par.add_run(a + ": ").bold = True
        par.add_run(b)
    if nph:
        runs(D.add_paragraph(), "[[RESULT: DRAFT, %d placeholders remain]]" % nph)
    D.add_heading("Summary of additional experiments", 2)
    for a in ADD:
        D.add_paragraph(re.sub(r"<[^>]+>", "", a).replace("&quot;", '"'), style="List Bullet")
    DOCX_SEEN = set()
    for k, p in blocks:
        if k == "h1":
            continue
        if k == "h2":
            D.add_heading(p, 1)
        elif k == "comment":
            cid0 = re.match(r"(\S+?)\.?\s", p + " ").group(1).rstrip(".")
            par0, _ = parent_id(cid0)
            t = D.add_table(rows=1, cols=1)
            c = t.rows[0].cells[0]
            shade(c, "E9EDF3")
            runs(c.paragraphs[0], "**" + p + "**")
            if par0 in DOCX_SEEN:
                continue
            DOCX_SEEN.add(par0)
            if par0 == "Ed-1":
                for g in ("Ed-0", "Ed-0b"):
                    if g in VCOM:
                        runs(c.add_paragraph(), VCOM[g])
            if par0 in VCOM:
                runs(c.add_paragraph(), "Reviewer comment, verbatim:")
                for part in VCOM[par0].split("\n"):
                    runs(c.add_paragraph(), part)
        elif k == "table":
            t = D.add_table(rows=len(p), cols=max(len(r) for r in p))
            t.style = "Table Grid"
            for ri, r in enumerate(p):
                for ci, v in enumerate(r):
                    runs(t.cell(ri, ci).paragraphs[0], ("**%s**" % v) if ri == 0 else v)
        elif k == "bullet":
            runs(D.add_paragraph(style="List Bullet"), p)
        elif k == "num":
            runs(D.add_paragraph(), p[0] + ". " + p[1])
        elif k == "p":
            runs(D.add_paragraph(), p)
    D.save(docx_path)
    print("DOCX:", docx_path)
except ImportError:
    print("python-docx not installed: .docx skipped (pip install python-docx)")
print("verbatim comments:", len(VCOM), "used:", len(VUSED), "NOT used:", sorted(set(VCOM) - VUSED))
print("draft entries unmatched to a parent:", UNMATCHED)
print("unresolved tex labels:", sorted(UNRESOLVED))
print("placeholders:", nph)
